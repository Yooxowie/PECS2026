import os
import docx
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_letters_doc():
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = DocInches(1.0)
        section.bottom_margin = DocInches(1.0)
        section.left_margin = DocInches(1.0)
        section.right_margin = DocInches(1.0)

    # Style definitions
    # Heading style
    def add_header_block(title, subtitle=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run1 = p.add_run("PHINMA UNIVERSITY OF PANGASINAN\n")
        run1.bold = True
        run1.font.size = DocPt(12)
        run1.font.color.rgb = DocRGBColor(27, 67, 50) # Dark Green
        
        run2 = p.add_run("COLLEGE OF HOSPITALITY AND TOURISM MANAGEMENT\n")
        run2.bold = True
        run2.font.size = DocPt(11)
        run2.font.color.rgb = DocRGBColor(212, 175, 55) # Harvest Gold
        
        run3 = p.add_run("BAM 205: Meetings, Incentives, Conferences, and Exhibitions (MICE)\n")
        run3.font.size = DocPt(9.5)
        run3.font.italic = True
        run3.font.color.rgb = DocRGBColor(100, 110, 105)

        if title:
            p_t = doc.add_paragraph()
            p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_t = p_t.add_run(title)
            run_t.bold = True
            run_t.font.size = DocPt(12.5)
            run_t.font.color.rgb = DocRGBColor(27, 67, 50)
            
        doc.add_paragraph().paragraph_format.space_after = DocPt(6)

    # -------------------------------------------------------------
    # LETTER 1: Gymnasium Reservation & Facility Access
    # -------------------------------------------------------------
    add_header_block("LETTER 1: REQUEST FOR GYMNASIUM RESERVATION & FACILITY ACCESS")

    p_date = doc.add_paragraph()
    r_date = p_date.add_run("Date: ")
    r_date.bold = True
    p_date.add_run("September 2, 2026")

    p_to = doc.add_paragraph()
    r_to = p_to.add_run("TO:       ")
    r_to.bold = True
    p_to.add_run("The Office of the Vice President for Academic Affairs /\nCampus Facilities & Athletics Office\nPHINMA University of Pangasinan\nArellano Street, Dagupan City")

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("SUBJECT:  ")
    r_sub.bold = True
    r_sub_t = p_sub.add_run("Venue Reservation, Technical Support, and Facility Access Request for PECS 2026 Conference")
    r_sub_t.bold = True

    doc.add_paragraph("Dear Administration,")
    
    doc.add_paragraph("Good day!")

    p_b1 = doc.add_paragraph(
        "The College of Hospitality and Tourism Management (CHTM) BAM 205 class respectfully requests permission to reserve and utilize the PHINMA-University of Pangasinan Gymnasium for the Pangasinan Eco-Haven & Culinary Showcase 2026 (PECS 2026) conference on September 17, 2026, from 7:00 AM to 5:00 PM (inclusive of setup, operational hours, and teardown)."
    )

    p_b2 = doc.add_paragraph(
        "We also humbly request authorization for student committee members to remain on campus beyond standard hours during setup on September 16, 2026 (until 7:00 PM), along with technical assistance for the following:"
    )

    reqs = [
        "1. Main stage projector / LED screen integration with HDMI connectivity.",
        "2. Sound system setup: 1 podium microphone, 4 wireless handheld microphones, and 2 lapel mics.",
        "3. Stage podium and chairs for panel discussions, as well as 500 spectator chairs arranged theater-style on the gym floor.",
        "4. Facility ventilation / cooling and lighting support throughout operational hours."
    ]
    for req in reqs:
        p_item = doc.add_paragraph(req)
        p_item.paragraph_format.left_indent = DocInches(0.4)
        p_item.paragraph_format.space_after = DocPt(3)

    doc.add_paragraph(
        "This conference serves as the final practical output for our BAM 205 course. Thank you very much for your continued guidance and support of student academic endeavors."
    )

    p_sign = doc.add_paragraph()
    p_sign.add_run("Respectfully yours,\n\n\n").font.size = DocPt(11)
    r_name = p_sign.add_run("[Your Name]\n")
    r_name.bold = True
    p_sign.add_run("Conference Director, PECS 2026\nCHTM — PHINMA University of Pangasinan")

    doc.add_page_break()

    # -------------------------------------------------------------
    # LETTER 2: Invitation Letter to Keynote Speaker
    # -------------------------------------------------------------
    add_header_block("LETTER 2: OFFICIAL INVITATION TO KEYNOTE SPEAKER")

    p_date2 = doc.add_paragraph()
    r_date2 = p_date2.add_run("Date: ")
    r_date2.bold = True
    p_date2.add_run("September 2, 2026")

    p_to2 = doc.add_paragraph()
    r_to2 = p_to2.add_run("TO:       ")
    r_to2.bold = True
    p_to2.add_run("[Speaker's Name / Professional Title]\n[Position / Designation]\n[Organization / Agency Name]")

    p_sub2 = doc.add_paragraph()
    r_sub2 = p_sub2.add_run("SUBJECT:  ")
    r_sub2.bold = True
    r_sub2_t = p_sub2.add_run("Invitation as Keynote Speaker – PECS 2026 Regional Conference")
    r_sub2_t.bold = True

    doc.add_paragraph("Dear [Speaker's Name],")
    doc.add_paragraph("Greetings from PHINMA University of Pangasinan!")

    doc.add_paragraph(
        "The College of Hospitality and Tourism Management cordially invites you to serve as a Keynote Speaker for the Pangasinan Eco-Haven & Culinary Showcase 2026 (PECS 2026) Conference on Thursday, September 17, 2026, at the PHINMA-UPang Gymnasium."
    )

    doc.add_paragraph(
        "Our conference theme, “Sustaining the Coastline: Strategies for Eco-Tourism Development and Heritage Culinary Preservation in Pangasinan,” aims to equip 500+ undergraduate delegates with actionable strategies for sustainable hospitality growth. Given your distinguished expertise in regional tourism frameworks and cultural preservation, we would be honored to have you deliver a 45-minute presentation followed by a 15-minute open Q&A session."
    )

    doc.add_paragraph(
        "Dedicated VIP hospitality, honorarium, tokens of appreciation, and campus parking privileges will be provided throughout the event. We look forward to confirming your esteemed participation."
    )

    p_sign2 = doc.add_paragraph()
    p_sign2.add_run("Warm regards,\n\n\n").font.size = DocPt(11)
    r_name2 = p_sign2.add_run("[Your Name]\n")
    r_name2.bold = True
    p_sign2.add_run("Speaker Relations Lead, PECS 2026\nCHTM — PHINMA University of Pangasinan")

    doc.add_page_break()

    # -------------------------------------------------------------
    # LETTER 3: Student Class Excuse Letter
    # -------------------------------------------------------------
    add_header_block("LETTER 3: STUDENT CLASS EXCUSE NOTIFICATION")

    p_date3 = doc.add_paragraph()
    r_date3 = p_date3.add_run("Date: ")
    r_date3.bold = True
    p_date3.add_run("September 14, 2026")

    p_to3 = doc.add_paragraph()
    r_to3 = p_to3.add_run("TO:       ")
    r_to3.bold = True
    p_to3.add_run("CHTM Faculty Members and Department Chairs\nPHINMA University of Pangasinan")

    p_sub3 = doc.add_paragraph()
    r_sub3 = p_sub3.add_run("SUBJECT:  ")
    r_sub3.bold = True
    r_sub3_t = p_sub3.add_run("Request for Class Excusal for PECS 2026 Student Organizers and Registered Delegates")
    r_sub3_t.bold = True

    doc.add_paragraph("Dear Professors,")

    doc.add_paragraph(
        "The organizing committee of PECS 2026 respectfully requests class excuse for participating BSHM and BSTM student organizers and registered delegates on Thursday, September 17, 2026, from 8:00 AM to 4:00 PM."
    )

    doc.add_paragraph(
        "These students will be actively managing conference logistics, main-stage operations, secretariat desks, and participating in session discussions at the PHINMA-UPang Gymnasium to fulfill BAM 205 (MICE) course requirements. Participating students remain accountable for all academic responsibilities and missed lectures."
    )

    doc.add_paragraph(
        "Thank you very much for your valued cooperation and continued support of our hospitality and tourism scholars."
    )

    p_sign3 = doc.add_paragraph()
    p_sign3.add_run("Sincerely,\n\n\n").font.size = DocPt(11)
    r_name3 = p_sign3.add_run("[Your Name]\n")
    r_name3.bold = True
    p_sign3.add_run("Conference Director, PECS 2026\nCHTM — PHINMA University of Pangasinan")

    doc.add_page_break()

    # -------------------------------------------------------------
    # FORM 4: Parent's / Guardian's Consent Form
    # -------------------------------------------------------------
    add_header_block("FORM 4: PARENT’S / GUARDIAN’S OFFICIAL CONSENT FORM")

    p_intro = doc.add_paragraph()
    p_intro.add_run("I, ______________________________________________________, parent / legal guardian of ")
    p_intro.add_run("______________________________________________________, student of BS [HM / TM] Year & Section __________, ")
    p_intro.add_run("hereby grant permission for my child to attend and participate in the official campus conference:")

    doc.add_paragraph().paragraph_format.space_after = DocPt(4)

    # Event details table box
    tbl = doc.add_table(rows=3, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [DocInches(2.2), DocInches(4.3)]
    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    details = [
        ("• Event Name:", "PECS 2026: Regional Conference on Sustainable Tourism & Heritage Culinary Innovation"),
        ("• Date & Operational Time:", "Thursday, September 17, 2026 | 7:00 AM – 5:00 PM\n(Including ingress, conference sessions, and egress)"),
        ("• Official Venue:", "PHINMA-University of Pangasinan Gymnasium, Arellano Street, Dagupan City")
    ]

    for idx, (label, val) in enumerate(details):
        c0 = tbl.cell(idx, 0)
        c1 = tbl.cell(idx, 1)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.color.rgb = DocRGBColor(27, 67, 50)
        
        p1 = c1.paragraphs[0]
        p1.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = DocPt(8)

    p_safe = doc.add_paragraph()
    r_safe = p_safe.add_run(
        "I understand that this event fulfills academic criteria for BAM 205 (Meetings, Incentives, Conferences, and Exhibitions - MICE) and that university safety and hall protocols will be strictly maintained by faculty supervisors."
    )
    r_safe.italic = True
    r_safe.font.size = DocPt(10)

    doc.add_paragraph().paragraph_format.space_after = DocPt(16)

    p_sig4 = doc.add_paragraph()
    p_sig4.add_run("___________________________________________________                Date: ________________________\n")
    r_sname = p_sig4.add_run("Parent / Guardian Signature over Printed Name\n\n")
    r_sname.bold = True
    p_sig4.add_run("Contact Number: ___________________________________")

    # Save to Downloads & Documents
    doc_paths = [
        r"C:\Users\Admin\Downloads\PECS_2026_Administrative_Letters.docx",
        r"C:\Users\Admin\Documents\PECS_2026_Administrative_Letters.docx",
        r"C:\Users\Admin\.gemini\antigravity\scratch\PECS_2026_Administrative_Letters.docx"
    ]
    for dp in doc_paths:
        try:
            doc.save(dp)
            print(f"Word Document saved to: {dp}")
        except Exception as e:
            print(f"Error saving doc to {dp}: {e}")

if __name__ == "__main__":
    create_letters_doc()
